import re, sys
from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

import os
MD = sys.argv[1] if len(sys.argv) > 1 else r"E:\Workbuddy\2026-07-27-11-58-27\manuscript\Fig5_draft.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else (os.path.splitext(MD)[0] + ".docx")

def add_runs(par, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for p in parts:
        if len(p) >= 4 and p.startswith('**') and p.endswith('**'):
            r = par.add_run(p[2:-2]); r.bold = True
        else:
            par.add_run(p)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)

lines = open(MD, encoding='utf-8').read().split('\n')
table_rows = []

def flush_table():
    if not table_rows:
        return
    ncol = len(table_rows[0])
    t = doc.add_table(rows=1, cols=ncol)
    try:
        t.style = 'Light Grid Accent 1'
    except Exception:
        pass
    for j, c in enumerate(table_rows[0]):
        cell = t.rows[0].cells[j]; cell.text = ''
        add_runs(cell.paragraphs[0], c.strip())
        for rr in cell.paragraphs[0].runs: rr.bold = True
    for row in table_rows[2:]:
        cells = t.add_row().cells
        for j, c in enumerate(row[:ncol]):
            cells[j].text = ''; add_runs(cells[j].paragraphs[0], c.strip())
    table_rows.clear()
    doc.add_paragraph('')

i = 0
while i < len(lines):
    line = lines[i]
    if line.strip().startswith('|'):
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if set(''.join(cells)) <= set('-: '):
            table_rows.append([''] * len(cells))
        else:
            table_rows.append(cells)
        i += 1; continue
    flush_table()
    s = line.strip()
    if s == '':
        i += 1; continue
    if s == '---':
        doc.add_paragraph(''); i += 1; continue
    if s.startswith('# '):
        add_runs(doc.add_heading(level=1), s[2:]); i += 1; continue
    if s.startswith('## '):
        add_runs(doc.add_heading(level=2), s[3:]); i += 1; continue
    if s.startswith('### '):
        add_runs(doc.add_heading(level=3), s[4:]); i += 1; continue
    if s.startswith('> '):
        p = doc.add_paragraph(); r = p.add_run(s[2:]); r.italic = True; r.font.size = Pt(9); i += 1; continue
    if s.startswith('![') and s.endswith(')'):
        m = re.match(r'^!\[(.*?)\]\((.*?)\)$', s)
        if m:
            cap, path = m.group(1), m.group(2)
            try:
                pimg = doc.add_paragraph(); pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = pimg.add_run()
                run.add_picture(path, width=Inches(6.0))
                shp = doc.inline_shapes[-1]
                if shp.height > Inches(9.0):
                    ratio = shp.width / shp.height
                    shp.height = Inches(9.0); shp.width = Emu(int(Inches(9.0) * ratio))
                pcap = doc.add_paragraph(); pcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rc = pcap.add_run(cap); rc.italic = True; rc.font.size = Pt(9)
            except Exception as e:
                doc.add_paragraph('[IMAGE MISSING: %s — %s]' % (path, e))
            i += 1; continue
    if s.startswith('- '):
        add_runs(doc.add_paragraph(style='List Bullet'), s[2:]); i += 1; continue
    add_runs(doc.add_paragraph(), s); i += 1
flush_table()
doc.save(OUT)
print("Saved", OUT)
